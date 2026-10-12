"""Knowledge retrieval: RAG persistence, access, upload limits and grounded citation contract."""
from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import NameObject, DictionaryObject, DecodedStreamObject
from fastapi.testclient import TestClient

from app.main import app
from app import embeddings, kb


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("SMARTBIZ_DB_PATH", str(tmp_path / "kb.sqlite3"))
    monkeypatch.setenv("ADMIN_API_KEY", "knowledge-admin-secret")
    monkeypatch.setenv("KB_SERVICE_KEY", "knowledge-service-secret")
    monkeypatch.setenv("RAG_EMBEDDING_MODE", "demo")
    monkeypatch.setenv("RAG_MIN_SIMILARITY", "0.25")
    monkeypatch.setenv("AGENT_MODE", "mock")
    with TestClient(app) as c:
        yield c


ADMIN = {"X-Admin-Key": "knowledge-admin-secret"}
SERVICE = {"X-KB-Service-Key": "knowledge-service-secret"}
CONTENT = "售后政策：工作日在线客服时间为09:00至18:00。" * 11


def upload(client, filename="售后规定.md", content=None, headers=None):
    return client.post(
        "/api/kb/documents", headers=headers or ADMIN,
        files={"file": (filename, content if content is not None else CONTENT.encode("utf8"), "text/plain")},
    )


def test_upload_list_search_delete_and_citations(client):
    assert upload(client, headers={"X-Admin-Key": "wrong"}).status_code == 403
    assert client.get("/api/kb/documents").status_code == 403
    assert client.get("/api/kb/search?q=客服").status_code == 403
    assert client.get("/api/kb/search?q=客服", headers={"X-KB-Service-Key": "invalid"}).status_code == 403
    response = upload(client)
    assert response.status_code == 201, response.text
    docid = response.json()["document_id"]
    assert response.json()["chunks"] >= 1
    assert len(client.get("/api/kb/documents", headers=ADMIN).json()) == 1
    matched = client.get("/api/kb/search?q=客服工作时间", headers=SERVICE)
    assert matched.status_code == 200
    items = matched.json()["sources"]
    assert items
    assert items[0]["document_id"] == docid
    assert items[0]["filename"] == "售后规定.md"
    assert items[0]["citation"] == "[1]"
    assert "在线客服" in items[0]["excerpt"]
    assert "similarity" in items[0]
    assert client.delete(f"/api/kb/documents/{docid}").status_code == 403
    assert client.delete(f"/api/kb/documents/{docid}", headers=ADMIN).status_code == 204
    assert client.get("/api/kb/search?q=客服工作时间", headers=SERVICE).json()["sources"] == []
    assert client.delete(f"/api/kb/documents/{docid}", headers=ADMIN).status_code == 404


def test_validation_limits_and_duplicate(client):
    assert upload(client, "not-valid.exe").status_code == 422
    assert upload(client, "empty.txt", b"").status_code == 422
    assert upload(client, "broken.md", b"\xff\xfe").status_code == 422
    assert upload(client, "oversize.txt", b"a" * (kb.MAX_UPLOAD_BYTES + 1)).status_code == 422
    assert upload(client).status_code == 201
    assert upload(client).status_code == 409
    assert client.get("/api/kb/search?q=x&limit=6", headers=SERVICE).status_code == 422
    assert client.get("/api/kb/search?q=", headers=SERVICE).status_code == 422
    assert client.get("/api/kb/search?q=" + "a" * 1001, headers=SERVICE).status_code == 422


def test_docx_extract_and_empty_pdf_error(client):
    doc = Document()
    doc.add_heading("企业知识库规范", level=1)
    doc.add_paragraph("文件管理要求：不允许上传未经授权的客户隐私文件。" * 12)
    fp = BytesIO()
    doc.save(fp)
    result = upload(client, "知识库.docx", fp.getvalue())
    assert result.status_code == 201, result.text
    assert upload(client, "image.pdf", b"not a real pdf").status_code == 422
    blank = PdfWriter()
    blank.add_blank_page(width=400, height=300)
    out = BytesIO()
    blank.write(out)
    assert upload(client, "scanned-like.pdf", out.getvalue()).status_code == 422


def test_pdf_with_embedded_text_keeps_page_number(client):
    writer = PdfWriter()
    page = writer.add_blank_page(width=500, height=500)
    font = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    font_ref = writer._add_object(font)
    page[NameObject("/Resources")] = DictionaryObject({
        NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_ref})
    })
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 50 400 Td (Support policy: weekdays 09:00 to 18:00.) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    fp = BytesIO()
    writer.write(fp)
    assert upload(client, "policy.pdf", fp.getvalue()).status_code == 201
    result = client.get("/api/kb/search?q=Support%20policy%20weekdays", headers=SERVICE)
    assert result.status_code == 200
    assert result.json()["sources"][0]["page"] == 1
    assert "Support policy" in result.json()["sources"][0]["excerpt"]


def test_kb_ask_requires_credentials_and_reports_real_sources(client, monkeypatch):
    upload(client)
    user = client.post("/api/sessions").json()
    request = {"question": "客服工作时间是几点？", "limit": 3}
    assert client.post("/api/kb/ask", headers=SERVICE, json=request).status_code == 401
    assert client.post("/api/kb/ask", headers={"X-Session-Token": user["session_token"]}, json=request).status_code == 403
    headers = {**SERVICE, "X-Session-Token": user["session_token"]}
    # Mock mode must not pretend it produced a real grounded LLM answer.
    assert client.post("/api/kb/ask", headers=headers, json=request).status_code == 503
    monkeypatch.setenv("AGENT_MODE", "langflow")
    seen = {}

    async def fake_llm(prompt, session_id):
        seen["prompt"], seen["session_id"] = prompt, session_id
        return "根据售后服务规定，工作日在线客服时间是 09:00 至 18:00。[1]"

    monkeypatch.setattr("app.main.ask_agent", fake_llm)
    result = client.post("/api/kb/ask", headers=headers, json=request)
    assert result.status_code == 200, result.text
    data = result.json()
    assert data["evidence_status"] == "retrieved_not_verified"
    assert data["sources"][0]["citation"] == "[1]"
    assert seen["session_id"] == "kb:" + user["session_id"]
    assert "<untrusted_retrieval>" in seen["prompt"]
    assert "客服工作时间是几点" in seen["prompt"]


def test_no_source_does_not_call_model(client, monkeypatch):
    session = client.post("/api/sessions").json()
    headers = {**SERVICE, "X-Session-Token": session["session_token"]}

    async def should_not_call(*args):
        raise AssertionError("must not invoke model without evidence")

    monkeypatch.setattr("app.main.ask_agent", should_not_call)
    result = client.post("/api/kb/ask", headers=headers, json={"question": "有什么资料？"})
    assert result.status_code == 200
    assert result.json()["evidence_status"] == "no_sources"
    assert result.json()["sources"] == []


def test_change_embedding_model_fails_closed(client, monkeypatch):
    assert upload(client).status_code == 201
    monkeypatch.setenv("RAG_EMBEDDING_MODE", "openai")
    monkeypatch.setenv("RAG_EMBEDDING_MODEL", "another-model")
    monkeypatch.setenv("RAG_EMBEDDING_API_KEY", "test-not-real")
    # Stub a real provider, only to test the index specification mismatch.
    async def fake_embed(texts):
        return [embeddings.demo_hash_embedding(t) for t in texts]
    monkeypatch.setattr("app.main.embeddings.embed_many", fake_embed)
    result = client.get("/api/kb/search?q=客服", headers=SERVICE)
    assert result.status_code == 409
    assert "重建索引" in result.json()["detail"]


def test_chunking_overlap_and_page_tracking():
    text = "午间客服规定。" * 170
    chunks = kb.chunk_sections([(5, text)])
    assert len(chunks) >= 2
    assert all(x["page"] == 5 for x in chunks)
    assert all(0 < len(x["content"]) <= kb.CHUNK_SIZE for x in chunks)
    assert chunks[0]["content"][-kb.CHUNK_OVERLAP:] in chunks[1]["content"]
    assert kb.split_text("\n\n") == []
    with pytest.raises(kb.KbInputError):
        kb.extract_sections("file.txt", b"")


def test_embedding_contract_guards():
    assert len(embeddings.demo_hash_embedding("你好世界")) == 256
    with pytest.raises(embeddings.EmbeddingError):
        embeddings.normalize([float("nan")])
    with pytest.raises(embeddings.EmbeddingError):
        embeddings.normalize([0.0] * 4)
    with pytest.raises(embeddings.EmbeddingError):
        embeddings.demo_hash_embedding("   ")


def test_admin_service_secrets_required(tmp_path, monkeypatch):
    monkeypatch.setenv("SMARTBIZ_DB_PATH", str(tmp_path / "kb2.sqlite3"))
    monkeypatch.delenv("ADMIN_API_KEY", raising=False)
    monkeypatch.delenv("KB_SERVICE_KEY", raising=False)
    with TestClient(app) as local:
        assert local.get("/api/kb/documents").status_code == 503
        assert local.get("/api/kb/search?q=客服").status_code == 503
