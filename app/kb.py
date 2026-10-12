"""Small, auditable RAG index backed by SQLite, for educational deployment only.

No original file download: only cleaned text chunks and vectors are persisted.
Do not upload confidential client files until authentication, tenant isolation,
retention, encrypted storage, and access policies are implemented.
"""
from array import array
import hashlib
from io import BytesIO
import math
from pathlib import Path
import re
import sqlite3
import uuid
from zipfile import ZipFile, BadZipFile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from lxml.etree import XMLSyntaxError
from pypdf import PdfReader
from pypdf.errors import PdfReadError, PdfStreamError

from . import db

MAX_UPLOAD_BYTES = 2 * 1024 * 1024
MAX_DOCUMENT_CHARS = 65_000
MAX_PAGES = 25
MAX_CHUNKS_PER_DOCUMENT = 150
MAX_TOTAL_CHUNKS = 2_000
MAX_DOCUMENTS = 50
CHUNK_SIZE = 500
CHUNK_OVERLAP = 80
ALLOWED = {".pdf", ".docx", ".txt", ".md"}


class KbInputError(Exception):
    pass


class KbConflictError(Exception):
    pass


def initialize() -> None:
    with db.connection() as conn:
        conn.executescript("""
        CREATE TABLE IF NOT EXISTS kb_documents (
            id TEXT PRIMARY KEY,
            filename TEXT NOT NULL,
            sha256 TEXT NOT NULL UNIQUE,
            size_bytes INTEGER NOT NULL,
            embedding_spec TEXT NOT NULL,
            chunk_count INTEGER NOT NULL,
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
        );
        CREATE TABLE IF NOT EXISTS kb_chunks (
            id TEXT PRIMARY KEY,
            document_id TEXT NOT NULL REFERENCES kb_documents(id) ON DELETE CASCADE,
            position INTEGER NOT NULL,
            page INTEGER,
            content TEXT NOT NULL,
            vector BLOB NOT NULL,
            dimension INTEGER NOT NULL,
            UNIQUE(document_id, position)
        );
        CREATE INDEX IF NOT EXISTS idx_kb_chunks_doc ON kb_chunks(document_id);
        """)


def clean_text(text: str) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_sections(filename: str, raw: bytes) -> list[tuple[int | None, str]]:
    extension = Path(filename).suffix.lower()
    if extension not in ALLOWED:
        raise KbInputError("仅支持 PDF、DOCX、TXT、MD 文件")
    if not raw or len(raw) > MAX_UPLOAD_BYTES:
        raise KbInputError("文档不能为空，且大小不得超过 2 MiB")
    try:
        if extension in (".txt", ".md"):
            sections = [(None, clean_text(raw.decode("utf-8-sig")))]
        elif extension == ".pdf":
            reader = PdfReader(BytesIO(raw), strict=False)
            if reader.is_encrypted:
                raise KbInputError("暂不支持加密 PDF")
            if len(reader.pages) > MAX_PAGES:
                raise KbInputError("PDF 最多允许 25 页")
            sections = [(i + 1, clean_text(page.extract_text() or "")) for i, page in enumerate(reader.pages)]
        else:
            # Reject suspicious archive expansion before python-docx reads it.
            with ZipFile(BytesIO(raw)) as archive:
                infos = archive.infolist()
                if len(infos) > 150 or sum(x.file_size for x in infos) > 15 * 1024 * 1024:
                    raise KbInputError("DOCX 解压体积或内部文件数超过限制")
            doc = Document(BytesIO(raw))
            blocks = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    blocks.append(" | ".join(cell.text for cell in row.cells))
            sections = [(None, clean_text("\n".join(blocks)))]
    except (UnicodeDecodeError, BadZipFile, PackageNotFoundError, XMLSyntaxError, PdfReadError, PdfStreamError, ValueError, KeyError, OSError) as exc:
        raise KbInputError("文档无法解析，请检查编码或文件格式") from exc
    # Reject scanned PDFs if no embedded text. This service does not run OCR.
    sections = [(page, text) for page, text in sections if text]
    if not sections:
        raise KbInputError("未提取到文本；扫描版 PDF 需要另行添加 OCR")
    if sum(len(text) for _, text in sections) > MAX_DOCUMENT_CHARS:
        raise KbInputError("可提取文本不得超过 65000 字符")
    return sections


def split_text(text: str, chunk_size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> list[str]:
    if not (10 <= overlap < chunk_size <= 2000):
        raise ValueError("Invalid chunk policy")
    text = clean_text(text)
    chunks = []
    start = 0
    while start < len(text):
        end = min(start + chunk_size, len(text))
        # Prefer ending on a paragraph or sentence near the tail, if possible.
        if end < len(text):
            pivot = max((text.rfind(sep, start + chunk_size // 2, end) for sep in ("\n", "。", "！", "？")), default=-1)
            if pivot > start + chunk_size // 2:
                end = pivot + 1
        part = text[start:end].strip()
        if part:
            chunks.append(part)
        if end >= len(text):
            break
        start = end - overlap
    return chunks


def chunk_sections(sections: list[tuple[int | None, str]]) -> list[dict]:
    chunks = []
    for page, text in sections:
        for part in split_text(text):
            chunks.append({"position": len(chunks), "page": page, "content": part})
    if not chunks or len(chunks) > MAX_CHUNKS_PER_DOCUMENT:
        raise KbInputError("分块数量超出限制，请拆分文档再上传")
    return chunks


def store_document(filename: str, raw: bytes, embedding_spec: str, chunks: list[dict], vectors: list[list[float]]) -> dict:
    if len(chunks) != len(vectors):
        raise ValueError("Chunk/vector count mismatch")
    if len({len(v) for v in vectors}) != 1:
        raise ValueError("Mismatched vector dimensions")
    document_id = str(uuid.uuid4())
    digest = hashlib.sha256(raw).hexdigest()
    # Keep only basename, on both Windows and Linux hosts.
    name = filename.replace("\\", "/").split("/")[-1][:120]
    try:
        with db.connection() as conn:
            if conn.execute("SELECT COUNT(*) AS n FROM kb_documents").fetchone()["n"] >= MAX_DOCUMENTS:
                raise KbConflictError("当前知识库最多 50 个文档")
            existing = conn.execute("SELECT embedding_spec FROM kb_documents LIMIT 1").fetchone()
            if existing and existing["embedding_spec"] != embedding_spec:
                raise KbConflictError("当前嵌入模型和已入库文档不一致，请重建索引")
            total = conn.execute("SELECT COUNT(*) AS n FROM kb_chunks").fetchone()["n"]
            if total + len(chunks) > MAX_TOTAL_CHUNKS:
                raise KbConflictError("当前知识库最多 2000 个文本块")
            conn.execute(
                "INSERT INTO kb_documents(id,filename,sha256,size_bytes,embedding_spec,chunk_count) VALUES(?,?,?,?,?,?)",
                (document_id, name, digest, len(raw), embedding_spec, len(chunks)),
            )
            for chunk, vector in zip(chunks, vectors):
                conn.execute(
                    "INSERT INTO kb_chunks(id,document_id,position,page,content,vector,dimension) VALUES(?,?,?,?,?,?,?)",
                    (str(uuid.uuid4()), document_id, chunk["position"], chunk["page"], chunk["content"],
                     array("f", vector).tobytes(), len(vector)),
                )
    except sqlite3.IntegrityError as exc:
        raise KbConflictError("相同文件已存在，不允许重复上传") from exc
    return {"document_id": document_id, "filename": name, "chunks": len(chunks), "embedding_spec": embedding_spec}


def ensure_index_compatible(embedding_spec: str) -> bool:
    """Check active embedder BEFORE making a billable remote embedding call.

    Returns whether the index holds any document.
    """
    with db.connection() as conn:
        specs = {r[0] for r in conn.execute("SELECT DISTINCT embedding_spec FROM kb_documents").fetchall()}
    if specs and specs != {embedding_spec}:
        raise KbConflictError("知识库嵌入模型已变更，需要重建索引")
    return bool(specs)


def already_uploaded(raw: bytes) -> bool:
    digest = hashlib.sha256(raw).hexdigest()
    with db.connection() as conn:
        return conn.execute("SELECT 1 FROM kb_documents WHERE sha256=?", (digest,)).fetchone() is not None


def list_documents() -> list[dict]:
    with db.connection() as conn:
        rows = conn.execute("SELECT id,filename,size_bytes,chunk_count,embedding_spec,created_at FROM kb_documents ORDER BY created_at DESC,id DESC").fetchall()
    return [dict(row) for row in rows]


def delete_document(document_id: str) -> bool:
    with db.connection() as conn:
        cursor = conn.execute("DELETE FROM kb_documents WHERE id=?", (document_id,))
        return cursor.rowcount > 0


def search(query_vector: list[float], embedding_spec: str, limit: int = 3, min_score: float = 0.25) -> list[dict]:
    """Exact cosine ranking in Python. All vectors normalized at embedding time."""
    with db.connection() as conn:
        specs = {r[0] for r in conn.execute("SELECT DISTINCT embedding_spec FROM kb_documents").fetchall()}
        if specs and specs != {embedding_spec}:
            raise KbConflictError("知识库嵌入模型已变更，需要重建索引")
        rows = conn.execute("""
            SELECT c.id,c.position,c.page,c.content,c.vector,c.dimension,
                   d.id AS document_id,d.filename
            FROM kb_chunks c JOIN kb_documents d ON d.id=c.document_id
            WHERE d.embedding_spec=?
        """, (embedding_spec,)).fetchall()
    matches = []
    for row in rows:
        if row["dimension"] != len(query_vector):
            raise KbConflictError("向量维度发生变化，需要重建索引")
        vector = array("f")
        vector.frombytes(row["vector"])
        score = sum(a * b for a, b in zip(query_vector, vector))
        if not math.isfinite(score) or score < min_score:
            continue
        matches.append({
            "chunk_id": row["id"], "document_id": row["document_id"],
            "filename": row["filename"], "page": row["page"],
            "position": row["position"], "excerpt": row["content"],
            "similarity": round(score, 4),
        })
    matches.sort(key=lambda x: (-x["similarity"], x["filename"], x["position"]))
    for index, match in enumerate(matches[:limit], 1):
        match["citation"] = f"[{index}]"
    return matches[:limit]


def index_status(active_spec: str) -> dict:
    """Index metadata; no embedding API calls."""
    with db.connection() as conn:
        specs = [row[0] for row in conn.execute(
            "SELECT DISTINCT embedding_spec FROM kb_documents ORDER BY embedding_spec"
        ).fetchall()]
        document_count = conn.execute("SELECT COUNT(*) FROM kb_documents").fetchone()[0]
        chunk_count = conn.execute("SELECT COUNT(*) FROM kb_chunks").fetchone()[0]
    return {
        "active_spec": active_spec, "stored_specs": specs,
        "compatible": not specs or specs == [active_spec],
        "requires_reindex": bool(specs and specs != [active_spec]),
        "document_count": document_count, "chunk_count": chunk_count,
    }


def read_chunks_for_reindex() -> list[dict]:
    """Read only existing cleaned text chunks; original files are NOT stored."""
    with db.connection() as conn:
        rows = conn.execute("SELECT id, content FROM kb_chunks ORDER BY document_id,position").fetchall()
    return [dict(r) for r in rows]


def replace_all_vectors(snapshot: list[dict], vectors: list[list[float]], spec: str) -> dict:
    """Atomic replace; reject concurrent uploads/removals or edits during API calls."""
    if len(snapshot) != len(vectors):
        raise KbConflictError("重建期间向量数量不一致")
    if not snapshot:
        return {"chunks_reindexed": 0, "embedding_spec": spec}
    if len({len(v) for v in vectors}) != 1 or not vectors[0]:
        raise KbInputError("重建的向量维度无效")
    with db.connection() as conn:
        conn.execute("BEGIN IMMEDIATE")
        current = [dict(r) for r in conn.execute(
            "SELECT id, content FROM kb_chunks ORDER BY document_id,position").fetchall()]
        if current != snapshot:
            raise KbConflictError("重建时文档发生变化；未修改原索引，请重试")
        for row, vector in zip(snapshot, vectors):
            conn.execute("UPDATE kb_chunks SET vector=?,dimension=? WHERE id=?",
                         (array("f", vector).tobytes(), len(vector), row["id"]))
        conn.execute("UPDATE kb_documents SET embedding_spec=?", (spec,))
    return {"chunks_reindexed": len(snapshot), "embedding_spec": spec}
